# -*- coding: utf-8 -*-
"""Render a .pptx to PDF without an Office suite.

Covers what this deck uses: text boxes (runs, wrapping, alignment, vertical
anchor, line spacing), pictures, and tables.  Helvetica carries Arial's metrics;
glyphs outside WinAnsi fall back to DejaVu Sans run by run.
"""
import io, sys
from pptx import Presentation
from pptx.util import Emu
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader
from PIL import Image

FD = "/usr/share/fonts/truetype/dejavu/"  # DejaVu fallback for non-WinAnsi glyphs
pdfmetrics.registerFont(TTFont('DJ',  FD + 'DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('DJB', FD + 'DejaVuSans-Bold.ttf'))
PT = 72.0
EMU_PT = 914400.0 / PT
DEF_L = DEF_R = 0.1 * PT
DEF_T = DEF_B = 0.05 * PT


def winansi(ch):
    try:
        ch.encode('cp1252')
        return True
    except Exception:
        return False


def fname(bold, ansi):
    if ansi:
        return 'Helvetica-Bold' if bold else 'Helvetica'
    return 'DJB' if bold else 'DJ'


def segs(text, bold):
    """Split text into (chunk, fontname) runs by encodability."""
    out, cur, curk = [], '', None
    for ch in text:
        k = winansi(ch)
        if k != curk and cur:
            out.append((cur, fname(bold, curk))); cur = ''
        curk = k; cur += ch
    if cur:
        out.append((cur, fname(bold, curk)))
    return out


def width(text, size, bold):
    return sum(pdfmetrics.stringWidth(c, f, size) for c, f in segs(text, bold))


def draw_line(c, x, y, tokens):
    """tokens: [(text, size, bold, rgb)] laid out left to right at baseline y."""
    for text, size, bold, col in tokens:
        for chunk, f in segs(text, bold):
            c.setFont(f, size)
            c.setFillColorRGB(*[v / 255.0 for v in col])
            c.drawString(x, y, chunk)
            x += pdfmetrics.stringWidth(chunk, f, size)


def run_color(r, default=(16, 45, 67)):
    try:
        if r.font.color is not None and r.font.color.type is not None:
            rgb = r.font.color.rgb
            return (rgb[0], rgb[1], rgb[2])
    except Exception:
        pass
    return default


def para_tokens(p, dsize, dbold, dcol):
    toks = []
    for r in p.runs:
        if not r.text:
            continue
        toks.append((r.text,
                     r.font.size.pt if r.font.size else dsize,
                     bool(r.font.bold) if r.font.bold is not None else dbold,
                     run_color(r, dcol)))
    return toks


def wrap_tokens(toks, maxw):
    """Greedy word wrap across formatted tokens. Returns list of token lines."""
    lines, cur, curw = [], [], 0.0
    for text, size, bold, col in toks:
        parts = text.split(' ')
        for i, wd in enumerate(parts):
            piece = (' ' if i else '') + wd
            if not piece:
                continue
            pw = width(piece, size, bold)
            if cur and curw + pw > maxw:
                lines.append(cur)
                piece = wd
                pw = width(piece, size, bold)
                cur, curw = [], 0.0
            if piece:
                cur.append((piece, size, bold, col)); curw += pw
    if cur:
        lines.append(cur)
    return lines or [[]]


def text_block(c, tf, L, T, W, H, page_h, default_size=18.0):
    ml = tf.margin_left  if tf.margin_left  is not None else Emu(int(DEF_L * EMU_PT))
    mr = tf.margin_right if tf.margin_right is not None else Emu(int(DEF_R * EMU_PT))
    mt = tf.margin_top   if tf.margin_top   is not None else Emu(int(DEF_T * EMU_PT))
    ml, mr, mt = ml / EMU_PT, mr / EMU_PT, mt / EMU_PT
    inner_w = W - ml - mr
    anchor = str(tf.vertical_anchor)

    laid = []           # (token_line, line_height, alignment)
    for p in tf.paragraphs:
        toks = para_tokens(p, default_size, False, (16, 45, 67))
        align = str(p.alignment) if p.alignment is not None else 'LEFT'
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.18
        if not toks:
            laid.append(([], default_size * ls, align)); continue
        for ln in wrap_tokens(toks, inner_w):
            h = max([t[1] for t in ln] or [default_size]) * ls
            laid.append((ln, h, align))

    total = sum(h for _, h, _ in laid)
    if 'MIDDLE' in anchor:
        y = T + mt + max(0.0, (H - mt - total) / 2.0)
    elif 'BOTTOM' in anchor:
        y = T + H - total
    else:
        y = T + mt

    for ln, h, align in laid:
        if ln:
            lw = sum(width(t[0], t[1], t[2]) for t in ln)
            if 'CENTER' in align:
                x = L + ml + (inner_w - lw) / 2.0
            elif 'RIGHT' in align:
                x = L + ml + inner_w - lw
            else:
                x = L + ml
            asc = max(t[1] for t in ln) * 0.80
            draw_line(c, x, page_h - (y + asc), ln)
        y += h


def draw_picture(c, sh, L, T, W, H, page_h):
    try:
        im = Image.open(io.BytesIO(sh.image.blob)).convert('RGBA')
    except Exception:
        return
    bg = Image.new('RGBA', im.size, (255, 255, 255, 255))
    im = Image.alpha_composite(bg, im).convert('RGB')
    c.drawImage(ImageReader(im), L, page_h - (T + H), W, H)


def draw_table(c, sh, L, T, page_h):
    tbl = sh.table
    cw = [col.width / EMU_PT for col in tbl.columns]
    rh = [row.height / EMU_PT for row in tbl.rows]
    y = T
    for ri, row in enumerate(tbl.rows):
        x = L
        for ci, cell in enumerate(row.cells):
            w, h = cw[ci], rh[ri]
            fill = None
            try:
                if cell.fill.type is not None and int(cell.fill.type) == 1:
                    rgb = cell.fill.fore_color.rgb
                    fill = (rgb[0] / 255.0, rgb[1] / 255.0, rgb[2] / 255.0)
            except Exception:
                pass
            if fill:
                c.setFillColorRGB(*fill)
                c.rect(x, page_h - (y + h), w, h, stroke=0, fill=1)
            c.setStrokeColorRGB(0.87, 0.89, 0.90); c.setLineWidth(0.5)
            c.rect(x, page_h - (y + h), w, h, stroke=1, fill=0)
            text_block(c, cell.text_frame, x, y, w, h, page_h, default_size=16.0)
            x += w
        y += rh[ri]


def render(src, dst):
    prs = Presentation(src)
    PW, PH = prs.slide_width / EMU_PT, prs.slide_height / EMU_PT
    c = canvas.Canvas(dst, pagesize=(PW, PH))
    for sl in prs.slides:
        c.setFillColorRGB(1, 1, 1); c.rect(0, 0, PW, PH, stroke=0, fill=1)
        for sh in sl.shapes:
            try:
                L, T = sh.left / EMU_PT, sh.top / EMU_PT
                W, H = sh.width / EMU_PT, sh.height / EMU_PT
            except Exception:
                continue
            if sh.shape_type == 13 or sh.__class__.__name__ == 'Picture':
                draw_picture(c, sh, L, T, W, H, PH)
            elif sh.has_table:
                draw_table(c, sh, L, T, PH)
            elif sh.has_text_frame:
                text_block(c, sh.text_frame, L, T, W, H, PH)
        c.showPage()
    c.save()
    print('wrote', dst)


if __name__ == '__main__':
    render(sys.argv[1], sys.argv[2])
